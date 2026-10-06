"""Excel / CSV reading and writing.

Score sheets need an email column and a score column; the mapping sheet needs
an email column and a name column (or First + Last name columns). Detection is
forgiving: it looks for an "email" header in the first rows of every sheet and
falls back to sniffing cell contents when a sheet has no header row.
Supported inputs: .xlsx / .xlsm (openpyxl) and .csv. Legacy binary .xls is
rejected with a "save as .xlsx" message: its only parser (xlrd) is unmaintained
and carries an advertising-clause license, so it is deliberately not a dependency.
"""
import csv
import io
import math
import re
import zipfile

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.xml import DEFUSEDXML

# openpyxl only hardens its XML parsing (entity expansion, external entities)
# when defusedxml is importable; refuse to parse untrusted workbooks without it.
if not DEFUSEDXML:
    raise ImportError("defusedxml is required for safe Excel parsing: pip install -r requirements.txt")

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+'-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)*\.[A-Za-z]{2,}")
NUMBER_RE = re.compile(r"^[-+]?(\d+\.?\d*|\.\d+)$")
HEADER_SCAN_ROWS = 25
SNIFF_ROWS = 200
MAX_ISSUES = 25
MAX_UNZIPPED_BYTES = 100 * 1024 * 1024   # an .xlsx is a zip; cap what it may expand to
MAX_ZIP_MEMBERS = 5000
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

EMAIL_HEADERS = [r"^e ?mail( ?id| address| addr)?$", r"^(official |work |employee |emp |user )?e ?mail", r"e ?mail", r"^mail( ?id)?$"]
SCORE_HEADERS = [r"^(total )?score$", r"score", r"^points?$", r"points?", r"marks?", r"^total$", r"result"]
NAME_HEADERS = [r"^(full|employee|emp|display|associate|staff) ?name$", r"^name$", r"^(full|employee|emp|display) name"]


class UserError(Exception):
    """A problem with the uploaded file or request that the admin can fix."""


def round2(x):
    return round(float(x), 2) + 0.0  # + 0.0 turns -0.0 into 0.0


def _norm(v):
    return re.sub(r"[\s_\-.:/]+", " ", "" if v is None else str(v)).strip().lower()


def _blank(v):
    return v is None or (isinstance(v, str) and not v.strip())


def _trunc(v, n=40):
    s = "" if v is None else str(v)
    return s if len(s) <= n else s[:n] + "..."


def clean_email(v):
    if _blank(v) or isinstance(v, (bool, int, float)):
        return None
    m = EMAIL_RE.search(str(v))
    return m.group(0).lower() if m else None


def to_number(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(v) else None
    if isinstance(v, str):
        s = re.sub(r"[\s,%]", "", v)
        return float(s) if NUMBER_RE.match(s) else None
    return None


def _clean_name(v):
    if _blank(v):
        return None
    s = re.sub(r"\s+", " ", str(v)).strip()
    return s[:120]


# ---------------------------------------------------------------- reading ---

def read_sheets(data: bytes):
    """Return [{'name': str, 'rows': [[cell, ...], ...]}] where rows[i] is Excel row i+1."""
    if data[:4] == b"PK\x03\x04":
        sheets = _read_xlsx(data)
    elif data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        raise UserError("Old .xls format is not supported. In Excel choose File > Save As > "
                        "Excel Workbook (.xlsx), then upload that file.")
    else:
        sheets = _read_csv(data)
    sheets = [s for s in sheets if any(not _blank(c) for row in s["rows"] for c in row)]
    if not sheets:
        raise UserError("The workbook has no data.")
    return sheets


def _check_zip(data):
    """Reject zip bombs before openpyxl inflates them. zipfile never inflates a
    member past its declared size, so summing declared sizes is a sound cap."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            members = z.infolist()
    except zipfile.BadZipFile:
        raise UserError("Could not open the Excel file. Save it as .xlsx and try again.")
    if len(members) > MAX_ZIP_MEMBERS or sum(m.file_size for m in members) > MAX_UNZIPPED_BYTES:
        raise UserError("The Excel file is too large once unpacked. Split it or remove unused sheets.")


def _read_xlsx(data):
    _check_zip(data)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception:
        raise UserError("Could not open the Excel file. Save it as .xlsx and try again.")
    try:
        return [{"name": ws.title, "rows": [list(r) for r in ws.iter_rows(min_row=1, values_only=True)]}
                for ws in wb.worksheets]
    finally:
        wb.close()


def _read_csv(data):
    text = None
    for enc in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None or "\x00" in text[:2000]:
        raise UserError("Unsupported file type. Upload an .xlsx or .csv file.")
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [[c if c.strip() else None for c in r] for r in csv.reader(io.StringIO(text), dialect)]
    return [{"name": "CSV", "rows": rows}]


# ------------------------------------------------------- column detection ---

def _pick(cells, patterns, exclude=()):
    for pattern in patterns:
        rx = re.compile(pattern)
        for i, c in enumerate(cells):
            if c and i not in exclude and rx.search(c):
                return i
    return -1


def _column_stats(rows, start, col):
    emails = nums = text = filled = 0
    for row in rows[start:start + SNIFF_ROWS]:
        v = row[col] if col < len(row) else None
        if _blank(v):
            continue
        filled += 1
        if clean_email(v):
            emails += 1
        elif to_number(v) is not None:
            nums += 1
        else:
            text += 1
    return {"emails": emails, "nums": nums, "text": text, "filled": filled}


def _best_column(rows, start, score_fn, exclude=()):
    width = max((len(r) for r in rows[start:start + SNIFF_ROWS]), default=0)
    best, best_score = -1, 0
    for c in range(width):
        if c in exclude:
            continue
        s = score_fn(_column_stats(rows, start, c))
        if s > best_score:
            best, best_score = c, s
    return best


def _locate(sheet, kind):
    rows = sheet["rows"]
    for i, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        if any(clean_email(c) for c in row):
            break  # data rows have started, so this sheet has no header row
        cells = [_norm(c) for c in row]
        email = _pick(cells, EMAIL_HEADERS)
        if email == -1:
            continue

        def header(idx):
            v = row[idx] if idx < len(row) else None
            return str(v).strip() if not _blank(v) else f"Column {get_column_letter(idx + 1)}"

        if kind == "score":
            score = _pick(cells, SCORE_HEADERS, [email])
            if score == -1:
                score = _best_column(rows, i + 1, lambda s: s["nums"], [email])
            if score == -1:
                raise UserError(f'Found the email column ("{header(email)}") on sheet "{sheet["name"]}" '
                                f'but no score column. Add a header named "Score".')
            return {"header_row": i, "email": email, "score": score,
                    "labels": {"email": header(email), "score": header(score)}}

        name = _pick(cells, NAME_HEADERS, [email])
        first = _pick(cells, [r"^first ?name$", r"first ?name|given ?name"], [email])
        last = _pick(cells, [r"^last ?name$", r"last ?name|surname|family ?name"], [email])
        if name == -1 and first != -1:
            label = f"{header(first)} + {header(last)}" if last != -1 else header(first)
            return {"header_row": i, "email": email, "first": first, "last": last,
                    "labels": {"email": header(email), "name": label}}
        if name == -1:
            name = _pick(cells, [r"name"], [email])
        if name == -1:
            name = _best_column(rows, i + 1, lambda s: s["text"], [email])
        if name == -1:
            raise UserError(f'Found the email column ("{header(email)}") on sheet "{sheet["name"]}" '
                            f'but no name column. Add a header named "Name".')
        return {"header_row": i, "email": email, "name": name,
                "labels": {"email": header(email), "name": header(name)}}

    # No header row: sniff columns by content.
    email = _best_column(rows, 0, lambda s: s["emails"] if s["emails"] >= max(1, s["filled"] * 0.6) else 0)
    if email == -1:
        return None
    letter = lambda c: f"Column {get_column_letter(c + 1)}"
    other = _best_column(rows, 0, (lambda s: s["nums"]) if kind == "score" else (lambda s: s["text"]), [email])
    if other == -1:
        return None
    key = "score" if kind == "score" else "name"
    return {"header_row": -1, "email": email, key: other, "labels": {"email": letter(email), key: letter(other)}}


def _locate_in_workbook(sheets, kind):
    first_error = None
    for sheet in sheets:
        try:
            cols = _locate(sheet, kind)
        except UserError as e:
            first_error = first_error or e
            continue
        if cols:
            return sheet, cols
    if first_error:
        raise first_error
    what = '"Email ID" and "Score"' if kind == "score" else '"Email ID" and "Name"'
    raise UserError(f"Could not find {what} columns in any sheet. Download the template from this box to see the expected layout.")


def _cell(row, idx):
    return row[idx] if 0 <= idx < len(row) else None


# ---------------------------------------------------------------- parsing ---

def parse_score_file(data: bytes):
    sheet, cols = _locate_in_workbook(read_sheets(data), "score")
    scores, issues = {}, []
    rows = skipped = duplicates = 0

    for r in range(cols["header_row"] + 1, len(sheet["rows"])):
        row = sheet["rows"][r]
        raw_email, raw_score = _cell(row, cols["email"]), _cell(row, cols["score"])
        if _blank(raw_email) and _blank(raw_score):
            continue
        email = clean_email(raw_email)
        if not email:
            skipped += 1
            if len(issues) < MAX_ISSUES:
                issues.append(f'Row {r + 1}: invalid email "{_trunc(raw_email)}"')
            continue
        score = to_number(raw_score)
        if score is None:
            skipped += 1
            if len(issues) < MAX_ISSUES:
                issues.append(f'Row {r + 1}: score "{_trunc(raw_score)}" for {email} is not a number')
            continue
        if email in scores:
            duplicates += 1
        scores[email] = round2(scores.get(email, 0.0) + score)
        rows += 1

    if not scores:
        hint = f" First problem: {issues[0]}" if issues else ""
        raise UserError(f'No valid email/score rows found on sheet "{sheet["name"]}".{hint}')
    return {
        "sheet": sheet["name"], "columns": cols["labels"], "scores": scores,
        "rows": rows, "skipped": skipped, "duplicates": duplicates, "issues": issues,
        "people": len(scores), "total": round2(sum(scores.values())),
    }


def parse_mapping_file(data: bytes):
    sheet, cols = _locate_in_workbook(read_sheets(data), "name")
    names, issues = {}, []
    rows = skipped = duplicates = 0

    for r in range(cols["header_row"] + 1, len(sheet["rows"])):
        row = sheet["rows"][r]
        raw_email = _cell(row, cols["email"])
        if "name" in cols:
            raw_name = _cell(row, cols["name"])
        else:
            parts = [_cell(row, cols["first"]), _cell(row, cols["last"]) if cols["last"] != -1 else None]
            raw_name = " ".join(str(p).strip() for p in parts if not _blank(p))
        if _blank(raw_email) and _blank(raw_name):
            continue
        email, name = clean_email(raw_email), _clean_name(raw_name)
        if not email or not name:
            skipped += 1
            if len(issues) < MAX_ISSUES:
                issues.append(f'Row {r + 1}: ' + (f'invalid email "{_trunc(raw_email)}"' if not email else f"missing name for {email}"))
            continue
        if email in names:
            duplicates += 1
        names[email] = name
        rows += 1

    if not names:
        hint = f" First problem: {issues[0]}" if issues else ""
        raise UserError(f'No valid email/name rows found on sheet "{sheet["name"]}".{hint}')
    return {"sheet": sheet["name"], "columns": cols["labels"], "names": names, "rows": rows,
            "skipped": skipped, "duplicates": duplicates, "issues": issues, "people": len(names)}


# ---------------------------------------------------------------- writing ---

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="0B1A33")


def _style_sheet(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font, c.fill = HEADER_FONT, HEADER_FILL
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"


def _neutralise_formulas(cells):
    """Spreadsheet formula injection guard (CWE-1236). openpyxl stores any string
    starting with "=" as a live formula, and spreadsheet apps evaluate cells
    starting with + - @ too. Names and file names come from uploaded files, so
    force such cells to plain text (quotePrefix shows them unchanged, as text)."""
    for c in cells:
        if isinstance(c.value, str) and c.value.startswith(FORMULA_PREFIXES):
            c.data_type = "s"
            c.quotePrefix = True


def build_workbook(sheets):
    """sheets: [(title, rows, widths)] where rows[0] is the header row. Returns xlsx bytes."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows, widths in sheets:
        ws = wb.create_sheet(title[:31])
        for row in rows:
            ws.append(row)
            _neutralise_formulas(ws[ws.max_row])
        _style_sheet(ws, widths)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
